"""Excel (.xlsx) report generation service using openpyxl.

Generates beautifully formatted, client-ready spreadsheets for:
- Job 1: Technical SEO Crawl Issues
- Job 2: GSC Ranking Anomalies
- Job 3: Product Metadata QA & AI Recommendations
- Job 4: Redirect Chain & Migration Validation
- Full Dashboard: Multi-sheet executive SEO performance workbook
"""

import io
from datetime import datetime, timezone
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from models import db, CrawlIssue, RankingAnomaly, ProductQAIssue, RedirectCheck, JobRun, Dataset, PendingTask

# Styling constants
HEADER_FILL = PatternFill(start_color="312E81", end_color="312E81", fill_type="solid")
HEADER_FONT = Font(name="Arial", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Arial", size=14, bold=True, color="1E1B4B")
SUBTITLE_FONT = Font(name="Arial", size=10, italic=True, color="475569")

CRITICAL_FILL = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
CRITICAL_FONT = Font(name="Arial", size=10, bold=True, color="991B1B")

WARNING_FILL = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
WARNING_FONT = Font(name="Arial", size=10, bold=True, color="92400E")

PASS_FILL = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
PASS_FONT = Font(name="Arial", size=10, bold=True, color="065F46")

REGULAR_FONT = Font(name="Arial", size=10)
URL_FONT = Font(name="Arial", size=10, color="2563EB", underline="single")

THIN_BORDER = Border(
    left=Side(style='thin', color='E2E8F0'),
    right=Side(style='thin', color='E2E8F0'),
    top=Side(style='thin', color='E2E8F0'),
    bottom=Side(style='thin', color='E2E8F0'),
)


def _apply_header_styles(ws, row_idx):
    """Apply styling to a table header row."""
    ws.row_dimensions[row_idx].height = 26
    for cell in ws[row_idx]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER


def _auto_fit_columns(ws, min_col=1, max_col=None, max_width_limit=60):
    """Auto-fit column widths based on content length."""
    if max_col is None:
        max_col = ws.max_column

    for col in range(min_col, max_col + 1):
        col_letter = get_column_letter(col)
        max_len = 0
        for cell in ws[col_letter]:
            if cell.row in (1, 2):  # Skip main title rows for width calculation
                continue
            val = str(cell.value or '')
            if '\n' in val:
                val = max(val.split('\n'), key=len)
            max_len = max(max_len, len(val))
        adjusted_width = min(max(max_len + 4, 12), max_width_limit)
        ws.column_dimensions[col_letter].width = adjusted_width


def export_job1_crawl_issues(run_id, client_name="NuroSparx"):
    """Export Job 1 Crawl issues to Excel."""
    job_run = JobRun.query.get_or_404(run_id)
    issues = CrawlIssue.query.filter_by(run_id=run_id).order_by(
        CrawlIssue.severity.asc(), CrawlIssue.url.asc()
    ).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Crawl Issues"
    ws.views.sheetView[0].showGridLines = True

    # Title block
    ws["A1"] = f"{client_name} — Technical SEO Crawl Issues"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Run Date: {job_run.run_date.strftime('%B %d, %Y %I:%M %p')} | Total Issues Found: {len(issues)}"
    ws["A2"].font = SUBTITLE_FONT

    headers = [
        "URL", "Page Template", "Issue Category", "Severity",
        "Priority Score", "Why It Matters", "Recommended Action",
        "Owner", "Status", "Detected Date"
    ]
    ws.append([])  # blank row 3
    ws.append(headers)  # row 4
    _apply_header_styles(ws, 4)

    for i in issues:
        row_values = [
            i.url,
            i.template or "standard",
            i.issue,
            i.severity,
            round(i.priority_score, 1) if i.priority_score else 0,
            i.why_it_matters or "",
            i.recommended_action or "",
            i.owner or "SEO Team",
            i.status.title() if i.status else "Open",
            i.detected_date.strftime('%Y-%m-%d') if i.detected_date else "",
        ]
        ws.append(row_values)
        current_row = ws.max_row
        ws.row_dimensions[current_row].height = 20

        for col_idx, cell in enumerate(ws[current_row], 1):
            cell.font = REGULAR_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center")

            # Highlight severity column (col 4)
            if col_idx == 4:
                sev = str(cell.value).upper()
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if "CRITICAL" in sev:
                    cell.fill = CRITICAL_FILL
                    cell.font = CRITICAL_FONT
                elif "HIGH" in sev:
                    cell.fill = WARNING_FILL
                    cell.font = WARNING_FONT
                elif "PASS" in sev or "LOW" in sev:
                    cell.fill = PASS_FILL
                    cell.font = PASS_FONT

    _auto_fit_columns(ws)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_job2_gsc_anomalies(run_id, client_name="NuroSparx"):
    """Export Job 2 GSC Ranking Anomalies to Excel."""
    job_run = JobRun.query.get_or_404(run_id)
    anomalies = db.session.query(RankingAnomaly).filter_by(run_id=run_id).order_by(
        RankingAnomaly.alert_level.asc(), RankingAnomaly.clicks_change_pct.asc()
    ).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "GSC Ranking Anomalies"
    ws.views.sheetView[0].showGridLines = True

    # Title block
    ws["A1"] = f"{client_name} — GSC Query Ranking Anomalies"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Run Date: {job_run.run_date.strftime('%B %d, %Y %I:%M %p')} | Total Anomalies Detected: {len(anomalies)}"
    ws["A2"].font = SUBTITLE_FONT

    headers = [
        "Search Query", "Alert Level", "Alert Reason", "Branded Query?",
        "Clicks (Previous)", "Clicks (Current)", "Clicks Change (%)",
        "Impressions (Previous)", "Impressions (Current)", "Impressions Change (%)",
        "Position (Previous)", "Position (Current)", "Position Change"
    ]
    ws.append([])  # row 3
    ws.append(headers)  # row 4
    _apply_header_styles(ws, 4)

    for a in anomalies:
        row_values = [
            a.query,
            a.alert_level,
            a.alert_reason or "",
            "Branded" if a.is_branded else "Non-Branded",
            a.clicks_previous,
            a.clicks_current,
            round(a.clicks_change_pct, 1),
            a.impressions_previous,
            a.impressions_current,
            round(a.impressions_change_pct, 1),
            round(a.position_previous, 1),
            round(a.position_current, 1),
            round(a.position_change, 1),
        ]
        ws.append(row_values)
        current_row = ws.max_row
        ws.row_dimensions[current_row].height = 20

        for col_idx, cell in enumerate(ws[current_row], 1):
            cell.font = REGULAR_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center")

            # Center alert level (col 2)
            if col_idx == 2:
                level = str(cell.value).upper()
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if "HIGH" in level:
                    cell.fill = CRITICAL_FILL
                    cell.font = CRITICAL_FONT
                elif "MEDIUM" in level or "WATCH" in level:
                    cell.fill = WARNING_FILL
                    cell.font = WARNING_FONT
                elif "IMPROVED" in level or "STABLE" in level:
                    cell.fill = PASS_FILL
                    cell.font = PASS_FONT

            # Number formatting
            if col_idx in (5, 6, 8, 9):
                cell.number_format = '#,##0'
            elif col_idx in (7, 10):
                cell.number_format = '0.0"%"'
            elif col_idx in (11, 12, 13):
                cell.number_format = '0.0'

    _auto_fit_columns(ws)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_job3_product_qa(run_id, client_name="NuroSparx"):
    """Export Job 3 Product Metadata QA to Excel."""
    job_run = JobRun.query.get_or_404(run_id)
    issues = ProductQAIssue.query.filter_by(run_id=run_id).order_by(
        ProductQAIssue.product_name.asc()
    ).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Product Metadata QA"
    ws.views.sheetView[0].showGridLines = True

    ws["A1"] = f"{client_name} — Product Metadata Quality Assurance & AI Suggestions"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Run Date: {job_run.run_date.strftime('%B %d, %Y %I:%M %p')} | Total Products Audited: {len(issues)}"
    ws["A2"].font = SUBTITLE_FONT

    headers = [
        "Product URL", "Product Name", "Template", "Issue Category",
        "Current Value", "AI Suggested Value", "AI Confidence",
        "Approval Status", "Approved By", "Approved At", "Notes"
    ]
    ws.append([])
    ws.append(headers)
    _apply_header_styles(ws, 4)

    for p in issues:
        row_values = [
            p.url,
            p.product_name or "",
            p.template or "product",
            p.issue_type,
            p.current_value or "",
            p.suggested_value or "",
            f"{int(p.ai_confidence * 100)}%" if p.ai_confidence else "—",
            p.status.title() if p.status else "Pending",
            p.approved_by or "",
            p.approved_at.strftime('%Y-%m-%d %H:%M') if p.approved_at else "",
            p.notes or "",
        ]
        ws.append(row_values)
        current_row = ws.max_row
        ws.row_dimensions[current_row].height = 22

        for col_idx, cell in enumerate(ws[current_row], 1):
            cell.font = REGULAR_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center")

            # Status highlight (col 8)
            if col_idx == 8:
                st = str(cell.value).upper()
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if "APPROVED" in st:
                    cell.fill = PASS_FILL
                    cell.font = PASS_FONT
                elif "PENDING" in st:
                    cell.fill = WARNING_FILL
                    cell.font = WARNING_FONT

    _auto_fit_columns(ws)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_job4_redirect_checks(run_id, client_name="NuroSparx"):
    """Export Job 4 Redirect validation checks to Excel."""
    job_run = JobRun.query.get_or_404(run_id)
    checks = RedirectCheck.query.filter_by(run_id=run_id).order_by(
        RedirectCheck.pass_fail.asc(), RedirectCheck.old_url.asc()
    ).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Redirect Validation"
    ws.views.sheetView[0].showGridLines = True

    ws["A1"] = f"{client_name} — Redirect Chain & Migration Validation"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Run Date: {job_run.run_date.strftime('%B %d, %Y %I:%M %p')} | Total Redirects Tested: {len(checks)}"
    ws["A2"].font = SUBTITLE_FONT

    headers = [
        "Old (Source) URL", "Expected Target URL", "Expected Status",
        "Actual Status", "Actual Final Destination URL", "Redirect Hops",
        "Validation Result", "Failure Details & Diagnostics"
    ]
    ws.append([])
    ws.append(headers)
    _apply_header_styles(ws, 4)

    for c in checks:
        row_values = [
            c.old_url,
            c.expected_url or "",
            c.expected_status or 301,
            c.actual_status or "",
            c.actual_final_url or "",
            c.redirect_hops,
            c.pass_fail,
            c.failure_reason or "Redirect resolved directly with 301.",
        ]
        ws.append(row_values)
        current_row = ws.max_row
        ws.row_dimensions[current_row].height = 20

        for col_idx, cell in enumerate(ws[current_row], 1):
            cell.font = REGULAR_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center")

            # Result highlight (col 7)
            if col_idx == 7:
                res = str(cell.value).upper()
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if "CRITICAL" in res or "FAIL" in res:
                    cell.fill = CRITICAL_FILL
                    cell.font = CRITICAL_FONT
                elif "WARNING" in res:
                    cell.fill = WARNING_FILL
                    cell.font = WARNING_FONT
                elif "PASS" in res:
                    cell.fill = PASS_FILL
                    cell.font = PASS_FONT

    _auto_fit_columns(ws)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_full_dashboard(dataset_id=None, client_name="NuroSparx"):
    """Export complete multi-tab SEO performance workbook."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove default sheet

    # 1. Overview Sheet
    ws_ov = wb.create_sheet(title="Executive Overview")
    ws_ov.views.sheetView[0].showGridLines = True
    ws_ov["A1"] = f"{client_name} — Weekly SEO Performance Automation Report"
    ws_ov["A1"].font = TITLE_FONT
    ws_ov["A2"] = f"Generated on: {datetime.now().strftime('%B %d, %Y %I:%M %p')} | Confidential Client Report"
    ws_ov["A2"].font = SUBTITLE_FONT

    ov_headers = ["Audit Module", "Status", "Total Issues", "Critical Alerts", "Execution Time (s)", "Last Run Date"]
    ws_ov.append([])
    ws_ov.append(ov_headers)
    _apply_header_styles(ws_ov, 4)

    job_names = {
        1: "Job 1: Technical SEO Crawl Monitor",
        2: "Job 2: GSC Ranking Anomaly Detection",
        3: "Job 3: Product Metadata QA & AI Optimization",
        4: "Job 4: Redirect Chain & Migration Validation"
    }

    latest_runs = {}
    for job_num in range(1, 5):
        query = JobRun.query.filter_by(job_number=job_num, status='completed')
        if dataset_id:
            query = query.filter_by(dataset_id=dataset_id)
        run = query.order_by(JobRun.run_date.desc()).first()
        latest_runs[job_num] = run
        if run:
            ws_ov.append([
                job_names[job_num],
                "Completed",
                run.issues_found,
                run.critical_count,
                round(run.duration_seconds, 2),
                run.run_date.strftime('%Y-%m-%d %H:%M')
            ])
        else:
            ws_ov.append([job_names[job_num], "Not Run Yet", 0, 0, 0, "—"])
        
        cr = ws_ov.max_row
        ws_ov.row_dimensions[cr].height = 22
        for cell in ws_ov[cr]:
            cell.font = REGULAR_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center")

    _auto_fit_columns(ws_ov)

    # 2. Add individual sheets for any completed runs
    # Job 1
    if latest_runs.get(1):
        run1 = latest_runs[1]
        issues1 = CrawlIssue.query.filter_by(run_id=run1.id).all()
        ws1 = wb.create_sheet(title="1-Crawl Issues")
        ws1.views.sheetView[0].showGridLines = True
        ws1.append(["URL", "Template", "Issue", "Severity", "Action", "Owner", "Status"])
        _apply_header_styles(ws1, 1)
        for i in issues1[:500]:
            ws1.append([i.url, i.template or "standard", i.issue, i.severity, i.recommended_action or "", i.owner or "SEO Team", i.status or "Open"])
        _auto_fit_columns(ws1)

    # Job 2
    if latest_runs.get(2):
        run2 = latest_runs[2]
        anom2 = db.session.query(RankingAnomaly).filter_by(run_id=run2.id).all()
        ws2 = wb.create_sheet(title="2-GSC Anomalies")
        ws2.views.sheetView[0].showGridLines = True
        ws2.append(["Query", "Alert", "Prior Clicks", "Curr Clicks", "Change %", "Prior Pos", "Curr Pos", "Pos Change"])
        _apply_header_styles(ws2, 1)
        for a in anom2:
            ws2.append([a.query, a.alert_level, a.clicks_previous, a.clicks_current, round(a.clicks_change_pct, 1), round(a.position_previous, 1), round(a.position_current, 1), round(a.position_change, 1)])
        _auto_fit_columns(ws2)

    # Job 3
    if latest_runs.get(3):
        run3 = latest_runs[3]
        prod3 = ProductQAIssue.query.filter_by(run_id=run3.id).all()
        ws3 = wb.create_sheet(title="3-Product QA")
        ws3.views.sheetView[0].showGridLines = True
        ws3.append(["Product Name", "URL", "Issue", "Current Value", "Suggested Value", "Status"])
        _apply_header_styles(ws3, 1)
        for p in prod3[:500]:
            ws3.append([p.product_name or "", p.url, p.issue_type, p.current_value or "", p.suggested_value or "", p.status or "Pending"])
        _auto_fit_columns(ws3)

    # Job 4
    if latest_runs.get(4):
        run4 = latest_runs[4]
        red4 = RedirectCheck.query.filter_by(run_id=run4.id).all()
        ws4 = wb.create_sheet(title="4-Redirects")
        ws4.views.sheetView[0].showGridLines = True
        ws4.append(["Old URL", "Expected URL", "Actual Status", "Final Destination", "Hops", "Result", "Reason"])
        _apply_header_styles(ws4, 1)
        for r in red4:
            ws4.append([r.old_url, r.expected_url, r.actual_status, r.actual_final_url, r.redirect_hops, r.pass_fail, r.failure_reason])
        _auto_fit_columns(ws4)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_task_issues(task_id):
    """Export task-specific issues for developer/SEO executive action.
    
    Args:
        task_id: PendingTask ID.
        
    Returns:
        io.BytesIO Excel workbook buffer.
    """
    task = PendingTask.query.get_or_404(task_id)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Task Action Items"
    ws.views.sheetView[0].showGridLines = True

    # Title block
    ws["A1"] = f"{task.client_name} — Task: {task.title}"
    ws["A1"].font = TITLE_FONT
    appr_str = f"Approved by: {task.approved_by or 'Admin'} ({task.approved_at.strftime('%b %d, %Y') if task.approved_at else '—'})"
    ws["A2"] = f"Category: {task.category} | Severity: {task.severity} | Status: {task.status.upper()} | {appr_str}"
    ws["A2"].font = SUBTITLE_FONT

    if task.job_number == 1:
        headers = ["URL", "Page Template", "Issue Category", "Severity", "Priority", "Why It Matters", "Recommended Fix", "Developer Status", "Developer Notes"]
        ws.append([])
        ws.append(headers)
        _apply_header_styles(ws, 4)

        issues = CrawlIssue.query.filter_by(run_id=task.run_id, issue=task.category).all()
        if not issues:
            issues = CrawlIssue.query.filter_by(issue=task.category).all()

        for idx, issue in enumerate(issues, start=5):
            ws.append([
                issue.url,
                issue.template or "",
                issue.issue,
                issue.severity,
                issue.priority_score,
                issue.why_it_matters or "",
                issue.recommended_action or "",
                "Pending Fix" if task.status == "pending" else task.status.capitalize(),
                task.developer_notes or "",
            ])
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=idx, column=col_idx)
                cell.border = THIN_BORDER
                cell.font = REGULAR_FONT
                if col_idx == 1:
                    cell.font = URL_FONT
                    cell.hyperlink = issue.url

    elif task.job_number == 3:
        headers = ["Product Name", "Product URL", "Issue Type", "Current Value", "Suggested AI Value", "Status", "Developer Notes"]
        ws.append([])
        ws.append(headers)
        _apply_header_styles(ws, 4)

        items = ProductQAIssue.query.filter_by(run_id=task.run_id, issue_type=task.category).all()
        if not items:
            items = ProductQAIssue.query.filter_by(issue_type=task.category).all()

        for idx, itm in enumerate(items, start=5):
            ws.append([
                itm.product_name or "",
                itm.url,
                itm.issue_type,
                itm.current_value or "",
                itm.suggested_value or "",
                itm.status or "Pending",
                task.developer_notes or "",
            ])
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=idx, column=col_idx)
                cell.border = THIN_BORDER
                cell.font = REGULAR_FONT

    elif task.job_number == 4:
        headers = ["Old URL (Redirect Source)", "Target URL", "Status Code", "Redirect Chain", "Final Destination", "Validation Result", "Developer Notes"]
        ws.append([])
        ws.append(headers)
        _apply_header_styles(ws, 4)

        checks = RedirectCheck.query.filter_by(run_id=task.run_id).all()
        if not checks:
            checks = RedirectCheck.query.all()
        if 'CRITICAL' in task.category.upper():
            checks = [c for c in checks if c.pass_fail == 'CRITICAL']
        elif 'WARNING' in task.category.upper():
            checks = [c for c in checks if c.pass_fail == 'WARNING']

        for idx, chk in enumerate(checks, start=5):
            ws.append([
                chk.old_url,
                chk.expected_url or "",
                chk.actual_status or chk.expected_status,
                f"{chk.redirect_hops} hops",
                chk.actual_final_url or "",
                chk.pass_fail,
                task.developer_notes or "",
            ])
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=idx, column=col_idx)
                cell.border = THIN_BORDER
                cell.font = REGULAR_FONT

    else:
        headers = ["Query / Item", "Alert Category", "Previous Value", "Current Value", "Change", "Status", "Developer Notes"]
        ws.append([])
        ws.append(headers)
        _apply_header_styles(ws, 4)

        anomalies = RankingAnomaly.query.filter_by(run_id=task.run_id).all()
        if not anomalies:
            anomalies = RankingAnomaly.query.all()
        if 'HIGH' in task.category.upper():
            anomalies = [a for a in anomalies if a.alert_level == 'HIGH']
        elif 'MEDIUM' in task.category.upper():
            anomalies = [a for a in anomalies if a.alert_level == 'MEDIUM']
        elif 'WATCH' in task.category.upper():
            anomalies = [a for a in anomalies if a.alert_level == 'WATCH']

        for idx, a in enumerate(anomalies, start=5):
            ws.append([
                a.query,
                a.alert_level,
                a.clicks_previous,
                a.clicks_current,
                f"{round(a.clicks_change_pct, 1)}%",
                task.status.capitalize(),
                task.developer_notes or "",
            ])
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=idx, column=col_idx)
                cell.border = THIN_BORDER
                cell.font = REGULAR_FONT

    _auto_fit_columns(ws)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf

